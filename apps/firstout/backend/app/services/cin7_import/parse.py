from __future__ import annotations

import csv
import datetime
import re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import NamedTuple, Optional

from app.services.vat import inc_to_ex

CURRENCIES = {"ZAR", "USD", "EUR", "GBP", "AUD", "CNY"}
# Cin7 numbers the tiers. The owner confirmed:
# 1 Stockist ex VAT, 2 Wholesale ex VAT, 3 Decorator ex VAT, 4 Retail inc VAT.
STOCKIST_TIER = 1
WHOLESALE_TIER = 2
DECORATOR_TIER = 3
RETAIL_TIER = 4
# (tier number, price-list name, whether the Cin7 figure already includes VAT)
NAMED_PRICE_TIERS = (
    (STOCKIST_TIER, "Cin7 Stockist", False),
    (WHOLESALE_TIER, "Cin7 Wholesale", False),
    (DECORATOR_TIER, "Cin7 Decorator", False),
    (RETAIL_TIER, "Cin7 Retail", True),
)
FILL_DOWN_KEYS = {"sku", "product", "productname", "productsku"}
BARCODE_LIMIT = 64


def header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.lower())


def clip(value: Optional[str], limit: int) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text[:limit]


def as_bool(value: str) -> bool:
    return value.strip().lower() in {"true", "yes", "1"}


class NamedSkuPrices(NamedTuple):
    stockist_ex_vat: Optional[Decimal]
    wholesale_ex_vat: Optional[Decimal]
    decorator_ex_vat: Optional[Decimal]
    retail_ex_vat: Optional[Decimal]


def tier_price_list_name(tier: int) -> str:
    for number, name, _includes_vat in NAMED_PRICE_TIERS:
        if number == tier:
            return name
    return f"Cin7 Tier {tier}"


def tier_amount_ex_vat(raw: Optional[Decimal], tier: int) -> Optional[Decimal]:
    """Positive Cin7 tier amount as ex VAT. Blank and zero are no price.

    Tiers 1-3 are already ex VAT. Tier 4 (retail) includes 15% VAT and is
    converted with the same half-up cent rounding as every other SKU price.
    """
    if raw is None or raw <= 0:
        return None
    for number, _name, includes_vat in NAMED_PRICE_TIERS:
        if number == tier and includes_vat:
            return inc_to_ex(raw)
    return raw


def named_sku_prices(row: dict[str, str]) -> NamedSkuPrices:
    amounts = {
        tier: tier_amount_ex_vat(money(row.get(f"PriceTier{tier}")), tier)
        for tier, _name, _includes_vat in NAMED_PRICE_TIERS
    }
    return NamedSkuPrices(
        stockist_ex_vat=amounts[STOCKIST_TIER],
        wholesale_ex_vat=amounts[WHOLESALE_TIER],
        decorator_ex_vat=amounts[DECORATOR_TIER],
        retail_ex_vat=amounts[RETAIL_TIER],
    )


def _odd_price_detail(amount: Decimal, includes_vat: bool) -> Optional[str]:
    """A positive stored price that is still not a usable tier amount.

    Amounts below 1 are placeholders (0.01, 0.87). Retail includes 15% VAT;
    a figure that does not survive inc-to-ex and back to the same cent is odd.
    """
    reasons: list[str] = []
    if amount < 1:
        reasons.append("price is below 1")
    if includes_vat:
        ex_vat = inc_to_ex(amount)
        restored = (ex_vat * Decimal("1.15")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if restored != amount:
            reasons.append("including VAT does not round-trip to the same cent")
    if not reasons:
        return None
    return "; ".join(reasons)


def price_warnings(row: dict[str, str]) -> list[dict[str, str]]:
    """Missing or odd named-tier prices. Does not change the stored amount."""
    code = row.get("ProductCode", "").strip()
    found: list[dict[str, str]] = []
    for tier, name, includes_vat in NAMED_PRICE_TIERS:
        amount = money(row.get(f"PriceTier{tier}"))
        if amount is None or amount <= 0:
            found.append(
                {
                    "kind": "missing_tier",
                    "product_code": code,
                    "tier": name,
                    "detail": f"PriceTier{tier} has no positive price",
                }
            )
            continue
        detail = _odd_price_detail(amount, includes_vat)
        if detail:
            found.append(
                {
                    "kind": "odd_price",
                    "product_code": code,
                    "tier": name,
                    "detail": detail,
                }
            )
    return found


def money(value: object, places: str = "0.01") -> Optional[Decimal]:
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    try:
        return Decimal(text).quantize(Decimal(places), rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None


def integral_qty(value: object) -> tuple[Optional[int], Optional[str]]:
    if value is None:
        return None, "blank quantity"
    text = str(value).strip()
    if text == "":
        return None, "blank quantity"
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None, "quantity is not a number"
    if number <= 0:
        return None, "quantity is not a positive integer"
    if number != number.to_integral_value():
        return None, "quantity is fractional"
    return int(number), None


def parse_date(value: object) -> Optional[datetime.date]:
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d-%b-%y", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def find_file(directory: Path, prefix: str, suffix: str) -> Optional[Path]:
    matches = sorted(
        path
        for path in directory.iterdir()
        if path.is_file()
        # Cin7 names report workbooks with spaces ("Sale Documents_ALLTIME.xlsx");
        # treat a space and an underscore as the same character.
        and path.name.replace(" ", "_").startswith(prefix.replace(" ", "_"))
        and path.name.lower().endswith(suffix)
    )
    if not matches:
        return None
    return matches[0]


def realign_status_row(header: list[str], row: list[str]) -> tuple[list[str], bool]:
    """Repair a Cin7 party row whose Status cell is actually the currency.

    A shifted export puts ZAR in Status and a non-currency in Currency. Insert
    Active into the Status position so the remaining cells line up.
    """
    cells = [cell.strip() for cell in row]
    while len(cells) > len(header) and cells[-1] == "":
        cells.pop()
    if len(cells) < len(header):
        cells.extend([""] * (len(header) - len(cells)))
    try:
        status_i = header.index("Status")
        currency_i = header.index("Currency")
    except ValueError:
        return cells[: len(header)], False
    status = cells[status_i].upper()
    currency = cells[currency_i].upper()
    if status in CURRENCIES and currency not in CURRENCIES:
        cells = cells[:status_i] + ["Active"] + cells[status_i:]
        return cells[: len(header)], True
    return cells[: len(header)], False


def fill_down(header: list[str], rows: list[list[str]]) -> tuple[list[list[str]], int]:
    """Copy SKU and product name onto following blank cells.

    Cin7 detail sheets repeat a product as a block: the first line has the SKU
    and name, and later lines leave those cells blank.
    """
    indexes = [index for index, name in enumerate(header) if header_key(name) in FILL_DOWN_KEYS]
    last: dict[int, str] = {}
    filled = 0
    for row in rows:
        for index in indexes:
            if row[index]:
                last[index] = row[index]
            elif last.get(index):
                row[index] = last[index]
                filled += 1
    return rows, filled


def assign_barcodes(
    pairs: list[tuple[str, str]],
    taken: set[str],
) -> dict[str, tuple[str, bool]]:
    """Give every SKU a unique barcode.

    The first product code (sorted) keeps a shared Cin7 barcode. Later sharers
    and blank barcodes get ``CIN7-{code}``.
    """
    ordered = sorted(pairs, key=lambda item: item[0])
    owners: dict[str, list[str]] = {}
    for code, raw in ordered:
        barcode = raw.strip()
        if barcode:
            owners.setdefault(barcode, []).append(code)
    claimed = set(taken)
    assigned: dict[str, tuple[str, bool]] = {}
    for code, raw in ordered:
        barcode = raw.strip()
        keep = bool(barcode) and owners.get(barcode, [None])[0] == code and barcode not in claimed
        if keep:
            chosen = barcode[:BARCODE_LIMIT]
            generated = False
        else:
            chosen = _generated_barcode(code, claimed)
            generated = True
        claimed.add(chosen)
        assigned[code] = (chosen, generated)
    return assigned


def generated_barcode_rows(
    pairs: list[tuple[str, str]],
    assigned: dict[str, tuple[str, bool]],
) -> list[dict[str, str]]:
    """SKUs ``assign_barcodes`` would rewrite, with blank or shared as the reason.

    Reads the assignment. Does not choose a different barcode.
    """
    last_raw: dict[str, str] = {}
    for code, raw in sorted(pairs, key=lambda item: item[0]):
        last_raw[code] = raw.strip()
    rows: list[dict[str, str]] = []
    for code in sorted(assigned):
        chosen, was_generated = assigned[code]
        if not was_generated:
            continue
        raw = last_raw.get(code, "")
        rows.append(
            {
                "product_code": code,
                "reason": "blank" if raw == "" else "shared",
                "generated_barcode": chosen,
            }
        )
    return rows


def _generated_barcode(code: str, claimed: set[str]) -> str:
    base = f"CIN7-{code}"
    candidate = base[:BARCODE_LIMIT]
    number = 2
    while candidate in claimed or not candidate:
        suffix = f"-{number}"
        candidate = base[: BARCODE_LIMIT - len(suffix)] + suffix
        number += 1
    return candidate


def read_csv_records(path: Path, realign: bool = False) -> tuple[list[dict[str, str]], int]:
    text = path.read_text(encoding="utf-8-sig")
    parsed = list(csv.reader(text.splitlines()))
    if not parsed:
        return [], 0
    header = [cell.strip() for cell in parsed[0]]
    records: list[dict[str, str]] = []
    realigned = 0
    for raw in parsed[1:]:
        if not any(cell.strip() for cell in raw):
            continue
        if realign:
            cells, shifted = realign_status_row(header, raw)
            realigned += int(shifted)
        else:
            cells = [cell.strip() for cell in raw]
            while len(cells) > len(header) and cells[-1] == "":
                cells.pop()
            if len(cells) < len(header):
                cells.extend([""] * (len(header) - len(cells)))
            cells = cells[: len(header)]
        records.append({header[index]: cells[index] for index in range(len(header))})
    return records, realigned


def read_xlsx_all(path: Path) -> tuple[list[dict[str, str]], int]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["All"]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator)
        header: list[str] = []
        for cell in header_row:
            if cell is None:
                break
            header.append(str(cell).strip())
        matrix: list[list[str]] = []
        for raw in iterator:
            values: list[str] = []
            for index in range(len(header)):
                cell = raw[index] if raw is not None and index < len(raw) else None
                values.append(_cell_text(cell))
            if any(values):
                matrix.append(values)
    finally:
        workbook.close()
    matrix, filled = fill_down(header, matrix)
    records = [{header[index]: row[index] for index in range(len(header))} for row in matrix]
    return records, filled


def _cell_text(cell: object) -> str:
    if cell is None:
        return ""
    if isinstance(cell, datetime.datetime):
        return cell.date().isoformat()
    if isinstance(cell, datetime.date):
        return cell.isoformat()
    return str(cell).strip()


def positive_int(value: str) -> Optional[int]:
    number = money(value, "1")
    if number is None or number < 1:
        return None
    if number != number.to_integral_value():
        return None
    return int(number)


def carton_count(value: str) -> int:
    parsed = positive_int(value)
    if parsed is None:
        return 1
    return parsed


def reorder_min(value: str) -> Optional[int]:
    return positive_int(value)
